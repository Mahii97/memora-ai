import uuid
import os
from datetime import datetime
from docx import Document as DocxDocument

from app.models.source_document import SourceDocument, SourcePage
from app.ingestion.cleaner import TextCleaner


class DocxParser:

    def parse(self, filepath):

        doc = DocxDocument(filepath)

        pages = []
        paragraph_number = 1

        # ---------- Normal paragraphs ----------
        for paragraph in doc.paragraphs:

            cleaned = TextCleaner.clean(paragraph.text)

            if cleaned:
                pages.append(
                    SourcePage(
                        location=f"Paragraph {paragraph_number}",
                        text=cleaned
                    )
                )
                paragraph_number += 1

        # ---------- Tables (NEW) ----------
        for table in doc.tables:

            for row in table.rows:

                cells = []

                for cell in row.cells:

                    value = TextCleaner.clean(cell.text)

                    if value:
                        cells.append(value)

                if cells:
                    row_text = " | ".join(cells)

                    pages.append(
                        SourcePage(
                            location=f"Table Row {paragraph_number}",
                            text=row_text
                        )
                    )

                    paragraph_number += 1

        return SourceDocument(
            id=str(uuid.uuid4()),
            source_type="docx",
            source_name=os.path.basename(filepath),
            author="Unknown",
            created_at=datetime.now(),
            pages=pages
        )